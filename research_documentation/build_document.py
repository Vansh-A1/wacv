"""Create the plain-language research report from its reviewed Markdown source."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
doc = Document()
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin, section.bottom_margin = Inches(.78), Inches(.72)
section.left_margin, section.right_margin = Inches(.84), Inches(.84)
section.footer_distance = Inches(.32)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet'):
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.widow_control = True
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(8)
for name, size, before, after in [('Title', 25, 0, 14), ('Heading 1', 18, 0, 12),
                                 ('Heading 2', 12, 10, 6)]:
    style = doc.styles[name]
    style.font.size = Pt(size)
    style.font.bold = name != 'Title'
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
doc.styles['List Bullet'].font.size = Pt(10.5)
doc.styles['List Bullet'].paragraph_format.space_after = Pt(8)
doc.core_properties.title = 'Image Quality Research Progress and Findings'
doc.core_properties.subject = 'Plain language account of the image quality research and current status'
doc.core_properties.author = 'Research project'
doc.core_properties.keywords = 'image quality, pristine energy, artifact energy, severity ranking, reference statistics'
for line in (HERE / 'Research_Work_Explained.md').read_text().splitlines():
    line = line.strip()
    if not line:
        continue
    if line == '<!-- PAGE -->':
        doc.add_page_break()
    elif line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], 'Heading 2')
    elif line.startswith('- '):
        doc.add_paragraph(line[2:], 'List Bullet')
    else:
        doc.add_paragraph(line)
footer = section.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run()
run.font.name = 'Calibri'
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run._r.addnext(field)
path = HERE / 'Image_Quality_Research_Progress_and_Findings.docx'
doc.save(path)
print(path)
print('Sections:', (HERE / 'Research_Work_Explained.md').read_text().count('<!-- PAGE -->') + 1)
