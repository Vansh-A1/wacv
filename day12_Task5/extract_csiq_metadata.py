"""Read the author's workbook without modifying it; retain published metadata/DMOS."""
import csv,hashlib,json,math
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parent
frozen=root/'frozen_configuration.json'
assert hashlib.file_digest(frozen.open('rb'),'sha256').hexdigest()==(root/'frozen_configuration.sha256').read_text().split()[0]
path=root/'data/csiq/csiq.DMOS.xlsx'
w=openpyxl.load_workbook(path,read_only=True,data_only=True)
s=w['all_by_image'];header=list(next(s.iter_rows(min_row=4,max_row=4,values_only=True)))
assert header[3:9]==['image','dst_idx','dst_type','dst_lev','dmos_std','dmos']
rows=[]
for i,row in enumerate(s.iter_rows(min_row=5,values_only=True),5):
 image,dst_idx,dst_type,dst_lev,dmos_std,dmos=row[3:9]
 if image is None:continue
 ref=str(int(image)) if isinstance(image,(int,float)) and float(image).is_integer() else str(image)
 assert isinstance(dmos,(int,float)) and math.isfinite(dmos),(i,dmos)
 assert isinstance(dst_lev,(int,float)) and math.isfinite(dst_lev),(i,dst_lev)
 rows.append({'ref_id':ref.lower(),'distortion_type_original':str(dst_type),'distortion_index':int(dst_idx),
    'severity':int(dst_lev),'DMOS':dmos,'DMOS_std':dmos_std,'source_sheet':'all_by_image','source_excel_row':i})
assert len(rows)==866
assert len({(r['ref_id'],r['distortion_index'],r['severity']) for r in rows})==866
# Cross-check published final DMOS values on the independent all_by_distortion sheet.
second={}
for row in w['all_by_distortion'].iter_rows(min_row=5,values_only=True):
 if row[5] is None:continue
 ref=str(int(row[5])) if isinstance(row[5],(int,float)) and float(row[5]).is_integer() else str(row[5])
 second[(ref.lower(),int(row[3]),int(row[6]))]=row[11]
assert len(second)==866
for r in rows:assert abs(second[(r['ref_id'],r['distortion_index'],r['severity'])]-r['DMOS'])<1e-15
out=root/'csiq/published_metadata.csv'
with out.open('w',newline='') as f:
 writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
report={'source_url':'https://s2.smu.edu/~eclarson/csiq/csiq.DMOS.xlsx',
 'source_sha256':hashlib.file_digest(path.open('rb'),'sha256').hexdigest(),
 'extraction':'all_by_image!D4:I870; rows 5:870 are published metadata and final DMOS',
 'crosscheck':'all_by_distortion final DMOS column L agrees for all 866 entries',
 'N':len(rows),'reference_count':len(set(r['ref_id'] for r in rows)),
 'distortion_types':sorted(set(r['distortion_type_original'] for r in rows)),
 'severity_source':'published dst_lev column G; never inferred from DMOS',
 'severity_complete':True,'original_workbook_modified':False,
 'unrelated_source_issues':'Other historical/intermediate sheets include #REF! cells. They were preserved and not used; all final DMOS cells used here are finite and cross-sheet consistent.'}
(root/'csiq/metadata_extraction_audit.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps(report,indent=2))
